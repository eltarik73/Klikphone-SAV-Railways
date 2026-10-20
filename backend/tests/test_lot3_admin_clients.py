"""Lot 3 (check-up) : export clients, création publique de client, validation du marketing.

- Export CSV / Excel de toute la base : élévation admin exigée, trace dans admin_audit_log,
  aucune formule (=, +, -, @, tabulation, retour chariot) et CSV lisible par Excel français.
- POST /api/clients sans session : limite de débit par IP inconnue, longueurs et carte Camby bornées.
- Marketing : dates impossibles et textes plus longs que les colonnes → 422 (et non 500) ;
  l'événement envoyé par l'écran Community Manager (champ « date ») est accepté.

Base simulée (aucune base réelle nécessaire).
"""

import csv
import io
from contextlib import contextmanager
from datetime import datetime
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook

from app.api import auth as auth_api
from app.api import clients as clients_api
from app.api import marketing as marketing_api
from app.api.auth import create_admin_token, create_token, decode_token
from app.api.public_documents import LimiteurDebit
from app.main import app


def _h(admin=False, user="Sophie"):
    tok = create_token("accueil", user)
    h = {"Authorization": f"Bearer {tok}"}
    if admin:
        h["X-Admin-Token"] = create_admin_token(decode_token(tok))
    return h


LIGNES = [
    {"id": 1, "nom": '=HYPERLINK("http://example.test","x")', "prenom": "+1+1", "telephone": "0600000101",
     "email": "a@example.test", "societe": "@SUM(1)", "carte_camby": 0,
     "date_creation": datetime(2026, 10, 1, 9, 0), "nb_tickets": 2},
    {"id": 2, "nom": "Élodie", "prenom": "-2", "telephone": "0600000102", "email": "\tb@example.test",
     "societe": "\rS", "carte_camby": 1, "date_creation": None, "nb_tickets": 0},
]


def _fausse_base(lignes=LIGNES):
    """get_cursor simulé : le SELECT d'export renvoie `lignes`, chaque requête est notée."""
    executees = []

    @contextmanager
    def gc():
        cur = MagicMock()
        cur.execute.side_effect = lambda sql, params=None: executees.append((sql, params))
        cur.fetchall.return_value = lignes
        cur.fetchone.return_value = {"id": 99}
        yield cur

    return gc, executees


# ═══ Cellules sûres ═══════════════════════════════════════════════

def test_cellule_sure_neutralise_les_formules():
    for v in ("=1+1", "+33600000000", "-2", "@SUM(1)", "\tx", "\rx"):
        assert clients_api._cellule_sure(v) == "'" + v
    for v in ("Dupont", "0600000101", "", 12, None, "a=b"):
        assert clients_api._cellule_sure(v) == v


# ═══ Export : élévation admin ═════════════════════════════════════

@pytest.mark.parametrize("route", ["/api/clients/export/csv", "/api/clients/export/excel"])
def test_export_exige_elevation_admin(client, route):
    gc, executees = _fausse_base()
    with patch.object(clients_api, "get_cursor", gc):
        assert client.get(route).status_code == 401
        r = client.get(route, headers=_h())
        assert r.status_code == 403
        assert r.json()["detail"]["code"] == "admin_requis"
    # Rien n'a été lu : pas d'export sans élévation
    assert not any("FROM clients" in sql for sql, _ in executees)


def test_export_csv_excel_francais_sans_formule_et_trace(client):
    gc, executees = _fausse_base()
    with patch.object(clients_api, "get_cursor", gc):
        r = client.get("/api/clients/export/csv", headers=_h(admin=True))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    brut = r.content
    assert brut.startswith("﻿".encode("utf-8"))     # BOM : accents lus par Excel
    lignes = list(csv.reader(io.StringIO(brut.decode("utf-8-sig")), delimiter=";"))
    assert lignes[0] == clients_api.EXPORT_ENTETES
    assert lignes[1][1] == '\'=HYPERLINK("http://example.test","x")'
    assert lignes[1][2] == "'+1+1" and lignes[1][5] == "'@SUM(1)"
    assert lignes[2][1] == "Élodie" and lignes[2][2] == "'-2"
    assert lignes[2][4] == "'\tb@example.test" and lignes[2][5] == "'\rS"
    assert lignes[2][7] == ""    # date absente : cellule vide, pas « None »
    traces = [p for sql, p in executees if "admin_audit_log" in sql]
    assert len(traces) == 1
    assert traces[0][0] == "Sophie" and traces[0][2] == "export_clients_csv" and traces[0][3] == "2 client(s)"
    # Aucune donnée personnelle dans la trace
    assert not any("Élodie" in str(x) or "0600000101" in str(x) for x in traces[0])


def test_export_excel_cellules_texte_jamais_formule(client):
    gc, executees = _fausse_base()
    with patch.object(clients_api, "get_cursor", gc):
        r = client.get("/api/clients/export/excel", headers=_h(admin=True))
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    cellules = [c for ligne in ws.iter_rows(min_row=2) for c in ligne]
    assert not any(c.data_type == "f" for c in cellules)
    assert ws["B2"].value == '\'=HYPERLINK("http://example.test","x")' and ws["B2"].data_type == "s"
    assert ws["C2"].value == "'+1+1" and ws["F2"].value == "'@SUM(1)"
    assert ws["A2"].value == 1 and ws["I2"].value == 2     # les nombres restent des nombres
    assert any("export_clients_excel" in str(p) for sql, p in executees if "admin_audit_log" in sql)


# ═══ Création publique de client ══════════════════════════════════

@pytest.fixture
def limiteur_neuf():
    auth_api.reset_essais()      # IP du TestClient inconnue (aucun jeton vu)
    with patch.object(clients_api, "_LIMITEUR_CREATION", LimiteurDebit(limite=10, fenetre=600)):
        yield
    auth_api.reset_essais()


def test_creation_publique_limitee_par_ip(client, limiteur_neuf):
    gc, _ = _fausse_base()
    with patch.object(clients_api, "get_cursor", gc), patch.object(clients_api, "find_client_by_phone", lambda cur, t: None):
        codes = [client.post("/api/clients", json={"nom": "Fictif", "telephone": f"06000099{i:02d}"}).status_code
                 for i in range(11)]
        assert set(codes[:10]) == {200}
        assert codes[10] == 429
        # Le personnel (jeton) n'est jamais limité
        r = client.post("/api/clients", headers=_h(), json={"nom": "Fictif", "telephone": "0600009950"})
        assert r.status_code != 429


def test_creation_poste_connu_non_limitee(client, limiteur_neuf):
    gc, _ = _fausse_base()
    auth_api._IP_CONNUES.noter("testclient")
    with patch.object(clients_api, "get_cursor", gc), patch.object(clients_api, "find_client_by_phone", lambda cur, t: None):
        codes = {client.post("/api/clients", json={"nom": "Fictif", "telephone": f"06000098{i:02d}"}).status_code
                 for i in range(12)}
    assert codes == {200}


@pytest.mark.parametrize("corps", [
    {"nom": "a" * 101, "telephone": "0600009960"},
    {"nom": "Fictif", "prenom": "p" * 101, "telephone": "0600009960"},
    {"nom": "Fictif", "societe": "s" * 100_000, "telephone": "0600009960"},
    {"nom": "Fictif", "telephone": "0600009960", "carte_camby": 10 ** 12},
    {"nom": "Fictif", "telephone": "0600009960", "carte_camby": -1},
    {"nom": "Fictif", "telephone": "0600009960", "carte_camby": 2},
])
def test_creation_bornes_422(client, limiteur_neuf, corps):
    gc, executees = _fausse_base()
    with patch.object(clients_api, "get_cursor", gc):
        assert client.post("/api/clients", json=corps).status_code == 422
    assert executees == []


def test_creation_valeurs_normales_acceptees(client, limiteur_neuf):
    gc, _ = _fausse_base()
    with patch.object(clients_api, "get_cursor", gc), patch.object(clients_api, "find_client_by_phone", lambda cur, t: None):
        r = client.post("/api/clients", json={"nom": "N" * 100, "telephone": "0600009961", "carte_camby": 1})
    assert r.status_code == 200


@pytest.mark.parametrize("corps", [{"carte_camby": -2147483649}, {"carte_camby": 5}, {"nom": "x" * 101}])
def test_modification_bornes_422(client, corps):
    gc, executees = _fausse_base()
    with patch.object(clients_api, "get_cursor", gc):
        assert client.patch("/api/clients/10", headers=_h(), json=corps).status_code == 422
    assert executees == []


# ═══ Marketing : 422 au lieu de 500 ═══════════════════════════════

@pytest.fixture
def marketing_admin():
    gc, executees = _fausse_base(lignes=[])
    app.dependency_overrides[marketing_api._require_admin_marketing] = lambda: {"sub": "Nour", "target": "accueil"}
    with patch.object(marketing_api, "get_cursor", gc), patch.object(marketing_api, "_ensure_tables", lambda: None):
        yield executees
    app.dependency_overrides.pop(marketing_api._require_admin_marketing, None)


@pytest.mark.parametrize("route,corps", [
    ("/api/marketing/calendrier", {"titre": "t", "type": "post", "date_evenement": "2026-13-45 99:99"}),
    ("/api/marketing/calendrier", {"titre": "t", "type": "post", "date_evenement": "2026-10-20", "couleur": "#7C3AED00"}),
    ("/api/marketing/calendrier", {"titre": "t" * 256, "type": "post", "date_evenement": "2026-10-20"}),
    ("/api/marketing/templates", {"nom": "n" * 256, "contenu_template": "x"}),
    ("/api/marketing/templates", {"nom": "n", "contenu_template": "x", "couleur": "#12345678"}),
    ("/api/marketing/posts", {"titre": "t" * 501, "contenu": "x", "plateforme": "instagram"}),
    ("/api/marketing/posts", {"titre": "t", "contenu": "x", "plateforme": "p" * 51}),
    ("/api/marketing/posts/1/programmer", {"date_programmee": "2026-99-99T10:00:00"}),
])
def test_marketing_entrees_invalides_422(client, marketing_admin, route, corps):
    r = client.post(route, json=corps)
    assert r.status_code == 422, r.text
    assert not any("INSERT" in sql or "UPDATE" in sql for sql, _ in marketing_admin)


def test_marketing_put_invalides_422(client, marketing_admin):
    assert client.put("/api/marketing/calendrier/1", json={"date_evenement": "2026-02-30"}).status_code == 422
    assert client.put("/api/marketing/templates/1", json={"nom": "n" * 256}).status_code == 422
    assert client.put("/api/marketing/posts/1", json={"statut": "s" * 31}).status_code == 422


def test_marketing_calendrier_date_impossible_en_filtre(client, marketing_admin):
    assert client.get("/api/marketing/calendrier?date_debut=2026-13-45").status_code == 422
    assert client.get("/api/marketing/calendrier?date_debut=2026-10-01&date_fin=2026-10-31").status_code == 200


def test_marketing_evenement_de_l_ecran_community_manager(client, marketing_admin):
    """CommunityManager envoie {titre, type, date, heure, couleur, date_heure} : la date est lue sous « date »."""
    r = client.post("/api/marketing/calendrier", json={
        "titre": "Promo", "type": "post", "date": "2026-10-20", "heure": "10:00", "couleur": "#7C3AED",
        "date_heure": "2026-10-20T10:00:00",
    })
    assert r.status_code == 200, r.text
    inserts = [p for sql, p in marketing_admin if "INSERT INTO calendrier_marketing" in sql]
    assert len(inserts) == 1 and str(inserts[0][3]) == "2026-10-20" and inserts[0][4] == "10:00"
